#!/usr/bin/env python3
import openpyxl
import sys
import os

def analyze_excel_structure(file_path):
    """分析Excel文件的结构"""
    try:
        print(f"正在分析Excel文件: {file_path}")
        
        # 加载工作簿
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        
        print(f"工作表名称: {ws.title}")
        print(f"最大行数: {ws.max_row}")
        print(f"最大列数: {ws.max_column}")
        
        # 查看前几行的数据
        print("\n=== 前10行数据 ===")
        for row in ws.iter_rows(min_row=1, max_row=10, max_col=20, values_only=True):
            print(row)
        
        # 查看表头行（假设第一行是表头）
        print("\n=== 表头行（第一行）===")
        headers = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=1, column=col).value
            headers.append(cell_value)
            print(f"列{col}: {cell_value}")
        
        # 分析表头中的数字模式
        print("\n=== 表头数字分析 ===")
        header_numbers = []
        for i, header in enumerate(headers, 1):
            if header and isinstance(header, str):
                # 尝试提取数字
                import re
                match = re.search(r'(\d{3})', header)
                if match:
                    num = match.group(1)
                    header_numbers.append((i, header, num))
                    print(f"列{i}: '{header}' -> 数字: {num}")
        
        print(f"\n找到 {len(header_numbers)} 个包含数字的表头")
        
        # 分析列位置模式
        if header_numbers:
            print("\n=== 列位置模式分析 ===")
            for col_idx, header, num in header_numbers:
                # 计算当前公式的预期列位置
                expected_col = (int(num) - 1) * 4 + 2
                print(f"表头 '{header}' (数字{num}):")
                print(f"  实际列位置: {col_idx}")
                print(f"  公式计算列位置: {expected_col}")
                print(f"  是否匹配: {'是' if col_idx == expected_col else '否'}")
        
        return ws, headers
        
    except Exception as e:
        print(f"分析Excel文件时出错: {e}")
        return None, None

if __name__ == "__main__":
    # 尝试多个可能的Excel文件路径
    excel_paths = [
        "/mnt/d/ExcelHander/src/sample.xlsx",
        "/mnt/d/ExcelHander/excel/副本5-19奥.xlsx",
        "/mnt/d/ExcelHander/excel/5-19奥.xlsx"
    ]
    
    for path in excel_paths:
        if os.path.exists(path):
            print(f"\n{'='*60}")
            analyze_excel_structure(path)
            break
    else:
        print("未找到Excel文件")